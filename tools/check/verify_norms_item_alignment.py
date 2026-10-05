#!/usr/bin/env python3
"""核对常模与量表包是否兼容——**按面向比对题集**，而不是按题号顺序。

为什么改口径
============
第一版我逐题比较「官方 Full# 第 k 题」与「转换后 order=k 的题」，结果
298/300 不一致。查下来是**题序本身不同**：官方计分表按域分组
（N1..N6, E1..E6, O1..O6, A1..A6, C1..C6 循环），而 OpenScales 的题序是
另一套排列。**题号顺序不同并不等于常模错了**——常模算的是「面向分」的分布，
而面向分 = 该面向 10 题的均分。所以真正的兼容条件是：

    官方计分表里每个面向的**题集** == 量表包里同一面向的**题集**

只要题集相同，两者算出的面向分就是同一个量，常模可直接使用。
本脚本按这个条件核对：先按面向比对题集，再（对能对上的题）逐题核对正文。
"""

from __future__ import annotations

import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from _root import find_root  # noqa: E402
ROOT = find_root()
TOOL = os.path.join(ROOT, "archive", "datasets", "osf-johnson", "scoring",
                    "IPIP-NEO-300_ScoringTool.xlsx")
INST = os.path.join(ROOT, "data", "instruments", "ipip-neo.json")


def norm(s: str) -> str:
    s = s.lower().replace("\u2019", "'").replace("\u2018", "'")
    s = re.sub(r"[^a-z0-9' ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def main() -> int:
    sys.path.insert(0, ROOT)
    from content.facet_definitions import FACETS  # noqa: E402
    code2id = {str(v.get("code", "")).upper(): k for k, v in FACETS.items()}

    from openpyxl import load_workbook
    ws = load_workbook(TOOL, data_only=True)[
        load_workbook(TOOL, data_only=True).sheetnames[0]]
    off_facets: dict[str, set[str]] = {}
    off_sign: dict[str, str] = {}
    for row in ws.iter_rows(values_only=True):
        if not (row and isinstance(row[0], (int, float)) and 1 <= int(row[0]) <= 300):
            continue
        sign = str(row[1] or "")
        m = re.search(r"([NEOAC]\d{1,2})", sign)
        if not m:
            continue
        fid = code2id.get(m.group(1).upper())
        if not fid:
            continue
        off_facets.setdefault(fid, set()).add(norm(str(row[4] or "")))
        off_sign[norm(str(row[4] or ""))] = "-" if sign.startswith("-") else "+"

    inst = json.load(open(INST, encoding="utf-8"))
    proj_facets: dict[str, set[str]] = {}
    proj_sign: dict[str, str] = {}
    for it in inst["items"]:
        idd = it["dimensions"][0]
        proj_facets.setdefault(idd, set()).add(norm(it["text"]))
        proj_sign[norm(it["text"])] = "-" if it.get("key", 1) < 0 else "+"

    print(f"官方面向 {len(off_facets)} 个，题集合计 "
          f"{sum(len(v) for v in off_facets.values())} 题")
    print(f"项目面向 {len(proj_facets)} 个，题集合计 "
          f"{sum(len(v) for v in proj_facets.values())} 题\n")

    ok, bad = [], []
    for fid in sorted(set(off_facets) | set(proj_facets)):
        a, b = off_facets.get(fid, set()), proj_facets.get(fid, set())
        if a and a == b:
            ok.append(fid)
        else:
            bad.append((fid, len(a), len(b), len(a & b)))

    print(f"面向题集**完全相同**：{len(ok)} 个")
    if bad:
        print(f"不一致：{len(bad)} 个")
        for fid, na, nb, inter in bad:
            print(f"  {fid:22} 官方 {na} 题 / 项目 {nb} 题 / 交集 {inter}")
    print()
    if not bad:
        print("✅ 每个面向的题集一致——常模可安全用于这套量表包")
    else:
        print("❌ 有面向题集不一致——常模不可直接使用")

    # 反向方向核对（只对两边都有的题）
    common = set(off_sign) & set(proj_sign)
    mism = [t for t in common if off_sign[t] != proj_sign[t]]
    print(f"\n反向方向核对：共有题 {len(common)} 道，方向不一致 {len(mism)} 道")
    for t in mism[:6]:
        print(f"  官方 {off_sign[t]} / 项目 {proj_sign[t]}：{t[:60]}")
    if not mism and common:
        print("✅ 共同题的方向完全一致")
    return 0 if (not bad and not mism) else 1


if __name__ == "__main__":
    sys.exit(main())
