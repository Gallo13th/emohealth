#!/usr/bin/env python3
"""用归档里的 Johnson (2005) IPIP-NEO-300 全量作答算常模。

数据与依据
==========
· 作答：``archive/datasets/osf-johnson/raw/IPIP300.dat``
  定宽 333 字符；题项在**列 33–332**（300 位，每位 1–5）。
  （定宽结构由 ``tools/survey/probe_ipip300_layout.py`` 逐列统计得出，不是猜的。）
· 题号→面向与计分方向：``archive/datasets/osf-johnson/scoring/
  IPIP-NEO-300_ScoringTool.xlsx`` 的 ``Full# | Sign | Facet`` 三列。
  ``Sign`` 形如 ``+N1`` / ``-N1``，负号即反向计分。
  **不采用**「假定数据列序 = 项目题序」的做法——那正是会把题配错面向的地方。

产出
====
``data/norms/ipip-neo-300.json``：逐面向与逐域的百分位表（1..99）+
均值/标准差 + 来源与样本量，供界面把原始分换算成百分位。

口径
====
· 面向分 = 该面向 10 题的**均分**（IPIP-NEO-300 每面向 10 题）
· 域分   = 该域 6 个面向分的**均分**
· 反向题按 ``6 - v`` 还原为同向后再计分
· 百分位用**经验分布**（排序后线性插值），不用正态假设
"""

from __future__ import annotations

import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from _root import find_root  # noqa: E402
ROOT = find_root()
DATA = os.path.join(ROOT, "archive", "datasets", "osf-johnson", "raw", "IPIP300.dat")
TOOL = os.path.join(ROOT, "archive", "datasets", "osf-johnson", "scoring",
                    "IPIP-NEO-300_ScoringTool.xlsx")
CANON = os.path.join(ROOT, "data", "ipip-neo-300.json")
OUT = os.path.join(ROOT, "data", "norms", "ipip-neo-300.json")

ITEM_START, ITEM_N = 33, 300


def read_official_map() -> tuple[list[str], list[str], list[str]]:
    """返回 (facet_by_fullindex, sign_by_fullindex, text_by_fullindex)。"""
    from openpyxl import load_workbook
    wb = load_workbook(TOOL, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
    facets, signs, texts = [""] * ITEM_N, [""] * ITEM_N, [""] * ITEM_N
    for r in rows:
        if not isinstance(r[0], (int, float)):
            continue
        full = int(r[0])
        if not (1 <= full <= ITEM_N):
            continue
        signs[full - 1] = str(r[1] or "")
        facets[full - 1] = str(r[3] or "")
        texts[full - 1] = str(r[4] or "")
    return facets, signs, texts


def main() -> int:
    facets, signs, texts = read_official_map()
    missing = [i + 1 for i, f in enumerate(facets) if not f or not signs[i]]
    if missing:
        print(f"✗ 官方计分表里有 {len(missing)} 题缺面向或方向：{missing[:10]}")
        return 1
    print(f"官方计分表：{ITEM_N} 题，面向 "
          f"{len(set(facets))} 个，反向题 {sum(1 for s in signs if s.startswith('-'))} 道")

    # 与项目题序核对：两者必须一致，否则说明我把两套东西混了
    canon = json.load(open(CANON, encoding="utf-8"))
    items = sorted(canon["items"], key=lambda x: x["order"])
    proj_facet = [it["facet"] for it in items]
    proj_key = [it["key"] for it in items]
    # 项目里的 facet id 与官方 Facet 名可能写法不同，只核**反向题集合**
    off_rev = {i + 1 for i, s in enumerate(signs) if s.startswith("-")}
    proj_rev = {it["order"] for it in items if it["key"] < 0}
    print(f"  反向题：官方 {len(off_rev)} 道，项目 {len(proj_rev)} 道，"
          f"交集 {len(off_rev & proj_rev)} 道")
    if off_rev != proj_rev:
        only_o = sorted(off_rev - proj_rev)[:8]
        only_p = sorted(proj_rev - off_rev)[:8]
        print(f"  ⚠ 反向题集合不一致：仅官方有 {only_o}，仅项目有 {only_p}")
    else:
        print("  ✓ 反向题集合完全一致（官方计分表与项目产物互证）")

    # 面向 → 题号。
    # ⚠️ 命名约定不同，而且**名字不可靠**：官方表里是 "Modesty"，
    #    项目 FACETS 里是 name_en="Modesty/Humility"（id=modesty_humility）。
    #    但官方 ``Sign`` 列形如 ``+A5`` / ``-N1``，后缀正是项目 FACETS 的 ``code``
    #    ——用 code 映射才是稳的。所以优先按 code，名字只作兜底。
    sys.path.insert(0, ROOT)
    from content.facet_definitions import FACETS  # noqa: E402
    code2id = {str(v.get("code", "")).upper(): k for k, v in FACETS.items()}
    name2id = {str(v.get("name_en", "")).lower(): k for k, v in FACETS.items()}

    def to_id(sign: str, official_name: str) -> str | None:
        m = re.search(r"([NEOAC]\d{1,2})", sign or "", re.I)
        if m and m.group(1).upper() in code2id:
            return code2id[m.group(1).upper()]
        return name2id.get(official_name.strip().lower())

    fo: dict[str, list[int]] = {}
    for i, s in enumerate(signs):
        fid = to_id(s, facets[i])
        if fid:
            fo.setdefault(fid, []).append(i)
    sizes = sorted(set(len(v) for v in fo.values()))
    print(f"  按官方 code 映射：{len(fo)} 个面向，每面向题数 {sizes}")
    if len(fo) != 30:
        print(f"  ⚠ 应为 30 个面向，实际 {len(fo)}")

    # 读作答（定宽切列）
    sums: dict[str, list[float]] = {f: [] for f in fo}
    n_rows = n_used = 0
    with open(DATA, encoding="latin-1") as fh:
        for ln in fh:
            n_rows += 1
            ln = ln.rstrip("\r\n")
            if len(ln) < ITEM_START + ITEM_N:
                continue
            seg = ln[ITEM_START:ITEM_START + ITEM_N]
            if not seg.isdigit():
                continue
            vals = [int(c) for c in seg]
            if any(v < 1 or v > 5 for v in vals):
                continue                      # IPIP-NEO 作答只有 1..5
            n_used += 1
            for f, idxs in fo.items():
                tot = 0.0
                for i in idxs:
                    v = vals[i]
                    tot += (6 - v) if signs[i].startswith("-") else v
                sums[f].append(tot / len(idxs))
    print(f"\n数据文件 {n_rows} 行，有效作答 {n_used} 条")
    if n_used < 1000:
        print("✗ 有效作答太少，定宽切列可能不对")
        return 1

    # 域分 = 该域 6 个面向分的均分（面向映射已在上面按官方 code 完成）
    dom_of = {f: FACETS[f]["domain"] for f in sums if f in FACETS}
    print(f"域映射：{len(dom_of)}/{len(sums)} 个面向有域归属")

    def pct_table(vals: list[float]) -> dict:
        xs = sorted(vals)
        n = len(xs)
        table = {}
        for p in range(1, 100):
            k = (p / 100) * (n - 1)
            lo = int(k)
            hi = min(lo + 1, n - 1)
            table[str(p)] = round(xs[lo] + (xs[hi] - xs[lo]) * (k - lo), 4)
        mean = sum(xs) / n
        var = sum((x - mean) ** 2 for x in xs) / (n - 1)
        return {"percentiles": table, "mean": round(mean, 4),
                "sd": round(var ** 0.5, 4), "min": round(xs[0], 4),
                "max": round(xs[-1], 4)}

    out = {
        "instrument": "IPIP-NEO-300",
        "source": "Johnson (2005) IPIP-NEO 公开数据集（OSF）",
        "source_files": {
            "responses": "archive/datasets/osf-johnson/raw/IPIP300.dat",
            "responses_sha256":
                "1bbb7189f5f0f2883bb4493f956cf639178c5c8ee95319fd3881f6d9d2f189a3",
            "scoring_key": "archive/datasets/osf-johnson/scoring/"
                           "IPIP-NEO-300_ScoringTool.xlsx",
            "layout": "定宽 333 字符，题项在列 33–332"
                      "（tools/survey/probe_ipip300_layout.py 逐列统计得出）",
        },
        "n": n_used,
        # ⚠️ 样本量必须说清：项目既有元数据写的是 "N=20,993"（那是
        #    validity/ipip20993.dat 的**有效样本**）。本文件用的是
        #    raw/IPIP300.dat 里**全部通过校验的作答**，条数更多。
        #    两个数字来源不同，不能互相替代，所以这里写明差异。
        "n_note": "本文件用 raw/IPIP300.dat 中通过校验的全部作答；"
                  "项目既有元数据提到的 N=20,993 来自 validity/ipip20993.dat，"
                  "是另一个（更小的）子集。两者不可互换引用。",
        "validation": {
            "reverse_items_official": len(off_rev),
            "reverse_items_project": len(proj_rev),
            "reverse_items_intersection": len(off_rev & proj_rev),
            "consistency": ("官方计分表与项目产物在反向题集合上完全一致"
                            if off_rev == proj_rev else "反向题集合不一致，需人工核查"),
            "facet_count": len(fo),
            "items_per_facet": sorted(set(len(v) for v in fo.values())),
            "records_scanned": n_rows,
        },
        "unit": "facet_mean_of_10_items",
        "domain_unit": "mean_of_6_facets",
        "reverse_transform": "6 - v",
        "percentile_method": "经验分布线性插值（1..99）",
        "scale_range": [1, 5],
        "facets": {f: pct_table(v) for f, v in sorted(sums.items())},
        # ⚠️ 常模按**面向 id**（项目 FACETS 的键，如 ``achievementstriving``）索引，
        #    而量表包与界面用的是**面向显示名**（如 ``Achievement-Striving``）。
        #    两者对不上会让界面查不到常模——表现为百分位列全是「—」，
        #    而且不报错（我出图核对时才发现）。所以显式带上名字桥接。
        "facet_names": {f: str(FACETS[f].get("name_en", f)) for f in sums
                        if f in FACETS},
        "domains": {},
        "license": "IPIP 为公有领域；Johnson (2005) 数据集随论文公开（OSF）",
    }
    # 域分：按面向分逐行求均分
    from collections import defaultdict
    dom_members = defaultdict(list)
    for f, d in dom_of.items():
        dom_members[d].append(f)
    for d, fs in sorted(dom_members.items()):
        m = min(len(sums[f]) for f in fs)
        dom_vals = [sum(sums[f][i] for f in fs) / len(fs) for i in range(m)]
        out["domains"][d] = pct_table(dom_vals)
        print(f"  域 {d:18} 面向 {len(fs)} 个，n={m}")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
        fh.write("\n")
    print(f"\n✓ 已写出 {os.path.relpath(OUT, ROOT)}"
          f"（{os.path.getsize(OUT) / 1024:.1f} KB）")
    print(f"  n={n_used}｜面向 {len(out['facets'])}｜域 {len(out['domains'])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
